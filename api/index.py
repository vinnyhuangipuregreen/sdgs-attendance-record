import os
import io
import csv
import copy
import re
from urllib.parse import quote
from typing import Literal
from starlette.concurrency import run_in_threadpool
from api.storage import AttendanceStore, ConcurrentUpdate, DATABASE_NAME
from api.sso import resolve_session, install_sso
from api.attendance import taiwan_now, record_view, apply_action, checkin_window
from api.courses import course_fields
from api.record_import import RecordImport, import_records
from pathlib import Path
from datetime import datetime
from typing import Optional
from fastapi import FastAPI, Request, Form, HTTPException, UploadFile, File
from fastapi.responses import HTMLResponse, JSONResponse, RedirectResponse, FileResponse, Response
from fastapi.templating import Jinja2Templates
import openpyxl

from fastapi.staticfiles import StaticFiles

app = FastAPI(title="BLIA佛光永續學院簽到系統 - iPure Green")

BASE_DIR = Path(__file__).resolve().parent.parent
TEMPLATES_DIR = BASE_DIR / "templates"
UPLOAD_DIR = BASE_DIR / "uploadfile"
templates = Jinja2Templates(directory=str(TEMPLATES_DIR))
app.mount("/static", StaticFiles(directory=str(BASE_DIR / "static")), name="static")

# 跨版本相容模板渲染函式 (同時相容新舊版 Starlette / FastAPI，徹底解決 500 Internal Server Error)
def render_template(template_name: str, context: dict, request: Request, status_code: int = 200):
    ctx = dict(context) if context else {}
    ctx["request"] = request
    try:
        # 新版 Starlette (0.28.0+)
        return templates.TemplateResponse(request=request, name=template_name, context=ctx, status_code=status_code)
    except TypeError:
        # 標準與舊版 Starlette (以 name, context 作為位置參數)
        return templates.TemplateResponse(template_name, ctx, status_code=status_code)

SECRET_KEY = os.getenv("SECRET_KEY", "")
if not SECRET_KEY:
    if os.getenv("VERCEL"):
        raise RuntimeError("請設定課程登錄的 SECRET_KEY")
    SECRET_KEY = "local-attendance-development-only"

# 資料庫模型
SEED_DATA = {
    "admins": [
        {
            "id": 1,
            "email": "service@ipuregreen.org",
            "password": "2026ipuregreen",
            "status": "active",
            "role": "superadmin"
        },
        {
            "id": 2,
            "email": "vinnyhuang.ipuregreen@gmail.com",
            "password": "123456",
            "status": "active",
            "role": "admin"
        }
    ],
    "students": [
        {
            "id": 1, 
            "name": "黃雅筠", 
            "email": "vinnyhuang.ipuregreen@gmail.com", 
            "email_pathdemy": "vinnyhuang168@gmail.com",
            "phone": "0912345678"
        }
    ],
    "courses": [
        {
            "id": 1,
            "title": "導論：自然永續的核心觀念與倫理基礎",
            "course_date": "2026/09/15",
            "course_time": "18:00-22:00",
            "open_time": "2026-09-15 18:00:00"
        },
        {
            "id": 2,
            "title": "健康一體（One Health）與系統思維",
            "course_date": "2026/09/22",
            "course_time": "18:30-21:30",
            "open_time": "2026-09-22 18:00:00"
        }
    ],
    "attendances": {
        "1_1": {"status": "SIGNED_IN", "signed_at": "2026-09-15 18:05:12"}
    }
}

store = None

@app.on_event("startup")
async def initialize_database():
    global store
    store = AttendanceStore()
    await run_in_threadpool(store.initialize, SEED_DATA)

@app.middleware("http")
async def persistent_course_data(request: Request, call_next):
    if request.url.path in {"/admin", "/admin/login"} and request.headers.get("x-platform-embedded") != "1":
        return RedirectResponse("https://sdgs-action-platform.vercel.app/#admin-attendance", status_code=302)
    if request.url.path in {"/", "/healthz", "/sso/login", "/auth/sso/callback", "/sso/logout", "/logout"} or request.url.path.startswith("/static/"):
        return await call_next(request)
    try:
        await run_in_threadpool(resolve_session, request, {"students":[],"courses":[],"attendances":{}}, SECRET_KEY)
    except HTTPException as error:
        return JSONResponse(status_code=error.status_code, content={"error": error.detail})
    year = getattr(request.state,"academic_year", {"id":1})
    year_id = int(year["id"])
    expected_year = request.headers.get("x-academic-year")
    if expected_year and expected_year != str(year_id):
        return JSONResponse(status_code=409, content={"error":"學年度已變更，請重新整理頁面後再操作。"})
    requested_year = request.query_params.get("academic_year_id")
    if requested_year and requested_year != str(year_id):
        try:
            requested_id = int(requested_year)
        except ValueError:
            return JSONResponse(status_code=400,content={"error":"學年度參數無效"})
        return RedirectResponse("/sso/login?academic_year_id=" + str(requested_id), status_code=302)
    try:
        version, data = await run_in_threadpool(store.load, year_id)
    except (RuntimeError, OSError):
        return JSONResponse(status_code=503, content={"error": "課程資料庫暫時無法連線，請稍後再試"})
    request.state.db = data
    request.state.db_version = version
    original = copy.deepcopy(data)
    from api.sso import link_student, sync_roster
    profile = getattr(request.state,"platform_profile",None)
    if profile and profile["account_type"] == "student":
        request.state.sso_student = link_student(data, profile)
    roster = getattr(request.state, "platform_roster", None)
    if roster is not None:
        sync_roster(data, roster)
    if request.headers.get("authorization", "").startswith("Bearer ") and (request.url.path.startswith("/admin") or request.url.path.startswith("/api/admin")):
        if not get_current_admin(request):
            return JSONResponse(status_code=403, content={"error": "需要管理員或講師權限"})
    if request.url.path.startswith("/api/admin/students") and request.method in {"POST", "PUT", "PATCH", "DELETE"}:
        if not get_current_admin(request):
            return JSONResponse(status_code=401, content={"error": "未授權"})
        return JSONResponse(status_code=409, content={"error": "帳號已統一管理，請至系統管理後台維護", "management_url": "https://sdgs-action-platform.vercel.app/#admin"})
    response = await call_next(request)
    if original is not None and response.status_code < 400 and data != original:
        try:
            await run_in_threadpool(store.save, version, data, year_id)
        except ConcurrentUpdate as error:
            return JSONResponse(status_code=409, content={"error": str(error)})
        except (RuntimeError, OSError):
            return JSONResponse(status_code=503, content={"error": "資料尚未儲存成功，請稍後重新送出"})
    return response

@app.get("/healthz")
async def healthz():
    try:
        await run_in_threadpool(store.load)
    except (RuntimeError, OSError):
        return JSONResponse(status_code=503, content={"status": "error", "database": DATABASE_NAME})
    return {"status": "ok", "database": DATABASE_NAME, "storage": "turso" if store.url else "sqlite"}

@app.post("/api/admin/academic-years/initialize")
async def initialize_academic_year(request: Request):
    if not get_current_admin(request):
        raise HTTPException(403,"需要管理權限")
    item = await request.json()
    year_id = int(item.get("id",0))
    if year_id < 2:
        raise HTTPException(400,"學年度編號無效")
    await run_in_threadpool(store.initialize_year, year_id)
    return {"success":True}

@app.get("/api/admin/academic-years/{year_id}/dependencies")
async def academic_year_dependencies(year_id: int, request: Request):
    if not get_current_admin(request):
        raise HTTPException(403,"需要管理權限")
    _, data = await run_in_threadpool(store.load, year_id)
    return {"[課程登錄]課程資料":len(data["courses"]), "[課程登錄]出席紀錄":len(data["attendances"]),
            "[課程登錄]學員映射資料":len(data["students"])}

# --- 認證輔助函式 ---
def get_current_student(request: Request):
    return getattr(request.state, "sso_student", None)


def get_current_admin(request: Request):
    return getattr(request.state, "sso_admin", None)

# ==========================================
# 前台學員路由
# ==========================================
@app.get("/", response_class=HTMLResponse)
async def login_page(request: Request):
    return RedirectResponse(url="/dashboard", status_code=302)

@app.get("/logout")
async def logout():
    return RedirectResponse(url="/sso/logout", status_code=302)

@app.get("/dashboard", response_class=HTMLResponse)
async def dashboard_page(request: Request):
    student = get_current_student(request)
    if not student:
        return RedirectResponse(url="/sso/login", status_code=302)
    return student_dashboard(request, student)


@app.get("/admin/students/{platform_id}/dashboard", response_class=HTMLResponse)
async def admin_student_dashboard(platform_id: int, request: Request):
    if not get_current_admin(request):
        raise HTTPException(403, "需要管理員或講師權限")
    student = next((s for s in request.state.db["students"] if s.get("platform_key") == f"student:{platform_id}"), None)
    if student is None:
        raise HTTPException(404, "此學年度查無此學員")
    return student_dashboard(request, student, read_only=True)


def student_dashboard(request, student, read_only=False):
    now = taiwan_now()
    courses_view = []
    for c in request.state.db["courses"]:
        record = request.state.db["attendances"].get(f"{student['id']}_{c['id']}")
        courses_view.append({**c, **record_view(c, record, now)})

    return render_template(
        "dashboard.html",
        {
            "student": student,
            "read_only": read_only,
            "stats": {
                "signed": sum(c["signed"] for c in courses_view),
                "leave": sum(c["status"] == "ON_LEAVE" or bool(c["leave_at"]) for c in courses_view),
                "makeup": sum(c["status"] == "MAKEUP_DONE" for c in courses_view),
                "absent": sum(c["status"] == "ABSENT" for c in courses_view)
            },
            "courses": courses_view
        },
        request
    )

@app.get("/api/auth/logout")
async def student_logout():
    return RedirectResponse(url="/sso/logout", status_code=302)

@app.post("/api/auth/login")
async def api_login(request: Request, email: str = Form(...), phone: str = Form(...)):
    return JSONResponse(status_code=410, content={"error": "請使用系統單一登入", "login_url": "/sso/login"})

@app.post("/api/attendance/checkin")
async def api_checkin(request: Request):
    student = get_current_student(request)
    if not student:
        raise HTTPException(status_code=401, detail="請先登入")
    data = await request.json()
    return apply_action(request.state.db, student, data.get("course_id"), "checkin", taiwan_now())

@app.post("/api/attendance/leave")
async def api_leave(request: Request):
    student = get_current_student(request)
    if not student:
        raise HTTPException(status_code=401, detail="請先登入")
    data = await request.json()
    return apply_action(request.state.db, student, data.get("course_id"), "leave", taiwan_now(), data.get("reason", ""))

@app.post("/api/attendance/makeup")
async def api_makeup(request: Request):
    student = get_current_student(request)
    if not student:
        raise HTTPException(status_code=401, detail="請先登入")
    data = await request.json()
    return apply_action(request.state.db, student, data.get("course_id"), "makeup", taiwan_now())

# ==========================================
# 後台管理路由
# ==========================================
@app.get("/admin/login", response_class=HTMLResponse)
async def admin_login_page(request: Request):
    if get_current_admin(request):
        return RedirectResponse(url="/admin", status_code=302)
    return RedirectResponse(url="/sso/login?admin=true", status_code=302)

@app.post("/api/admin/auth/login")
async def admin_login(request: Request, email: str = Form(...), password: str = Form(...)):
    return JSONResponse(status_code=410, content={"error": "請使用系統單一登入", "login_url": "/sso/login?admin=true"})

@app.get("/api/admin/auth/logout")
async def admin_logout():
    return RedirectResponse(url="/sso/logout", status_code=302)

@app.get("/admin", response_class=HTMLResponse)
async def admin_page(request: Request):
    current_admin = get_current_admin(request)
    if not current_admin:
        return RedirectResponse(url="/admin/login", status_code=302)
    return render_template(
        "admin.html",
        {
            "admins": [current_admin],
            "students": request.state.db["students"],
            "courses": request.state.db["courses"],
            "course_counts": course_record_counts(request.state.db),
            "current_admin": current_admin,
            "checkin_window": checkin_window,
            "time_options": [f"{hour:02d}:{minute:02d}" for hour in range(24) for minute in (0, 30)]
        },
        request
    )


@app.get("/api/admin/courses/catalog")
async def course_catalog(request: Request):
    if not get_current_admin(request):
        raise HTTPException(403, "需要管理員權限")
    return [{"id": c["id"], "title": c["title"]} for c in request.state.db["courses"]]


def course_record_counts(data):
    counts = {course["id"]: {"SIGNED_IN": 0, "ON_LEAVE": 0, "MAKEUP_DONE": 0} for course in data["courses"]}
    for key, record in data["attendances"].items():
        course_id = int(key.rsplit("_", 1)[1])
        status = record.get("status")
        if course_id in counts and status in counts[course_id]:
            counts[course_id][status] += 1
            if status != "SIGNED_IN" and record.get("signed_at"):
                counts[course_id]["SIGNED_IN"] += 1
            if status == "MAKEUP_DONE" and record.get("leave_at"):
                counts[course_id]["ON_LEAVE"] += 1
    return counts


@app.get("/api/admin/attendance/export")
async def export_attendance_records(request: Request):
    if not get_current_admin(request):
        raise HTTPException(403, "需要管理權限")
    return {"version": request.state.db_version, "academic_year": request.state.academic_year,
            **{key: request.state.db[key] for key in ("students", "courses", "attendances")}}


@app.post("/api/admin/attendance/import")
async def import_attendance_records(request: Request, payload: RecordImport):
    admin = get_current_admin(request)
    if not admin:
        raise HTTPException(403, "需要管理權限")
    if payload.version != request.state.db_version:
        raise HTTPException(409, "資料已更新，請重新備份並核對後再匯入")
    count = import_records(request.state.db, payload.records, admin["email"])
    return {"success": True, "count": count, "course_counts": course_record_counts(request.state.db)}

# ==========================================
# 1. 學員名單管理：範本下載 & Excel 匯入校驗
# ==========================================
@app.get("/api/admin/students/sample-excel")
async def download_student_sample_excel(request: Request):
    if not get_current_admin(request):
        raise HTTPException(status_code=401, detail="未授權")
    
    UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
    sample_file = UPLOAD_DIR / "student_list_upload.xlsx"
    
    if not sample_file.exists():
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "學員名單"
        ws.append(["學員姓名", "Email", "Email(補課平台)", "聯絡電話"])
        ws.append(["黃雅筠", "vinnyhuang.ipuregreen@gmail.com", "vinnyhuang168@gmail.com", "0912345678"])
        ws.append(["林小明", "ming@example.com", "ming@example.com", "0987654321"])
        wb.save(str(sample_file))
        
    return FileResponse(
        path=str(sample_file),
        filename="student_list_upload.xlsx",
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

@app.post("/api/admin/students/import")
async def import_students(request: Request, file: UploadFile = File(...)):
    if not get_current_admin(request):
        raise HTTPException(status_code=401, detail="未授權")
    
    contents = await file.read()
    filename = file.filename.lower()

    if not (filename.endswith(".xlsx") or filename.endswith(".xls")):
        return JSONResponse(status_code=400, content={"error": "上傳失敗！僅支援 Excel 檔案格式 (.xlsx 或 .xls)"})

    try:
        wb = openpyxl.load_workbook(io.BytesIO(contents), data_only=True)
        sheet = wb.active
        rows = list(sheet.iter_rows(values_only=True)) if sheet else []
    except Exception as e:
        return JSONResponse(status_code=400, content={"error": f"Excel 檔案讀取失敗：{str(e)}"})

    if not rows or len(rows) < 1:
        return JSONResponse(status_code=400, content={"error": "上傳的 Excel 檔案為空，請確認內容！"})

    sample_file = UPLOAD_DIR / "student_list_upload.xlsx"
    expected_headers = ["學員姓名", "Email", "Email(補課平台)", "聯絡電話"]
    if sample_file.exists():
        try:
            s_wb = openpyxl.load_workbook(str(sample_file), data_only=True)
            s_ws = s_wb.active
            s_row = [str(cell.value).strip() for cell in s_ws[1] if cell.value is not None and str(cell.value).strip()]
            if s_row:
                expected_headers = s_row
        except:
            pass

    uploaded_headers = [str(col).strip() for col in rows[0] if col is not None and str(col).strip()]

    missing_fields = [h for h in expected_headers if h not in uploaded_headers]
    extra_fields = [h for h in uploaded_headers if h not in expected_headers]

    if missing_fields or extra_fields:
        err_msg = "匯入失敗！欄位格式不符合範例檔案 (student_list_upload.xlsx)：\n"
        if missing_fields:
            err_msg += f"❌ 缺少必要欄位：【{', '.join(missing_fields)}】\n"
        if extra_fields:
            err_msg += f"⚠️ 多出未定義欄位：【{', '.join(extra_fields)}】\n"
        err_msg += f"👉 標準欄位格式為：【{', '.join(expected_headers)}】\n請點擊旁邊「📥 下載範例Excel」核對格式後重新上傳。"
        return JSONResponse(status_code=400, content={"error": err_msg})

    name_idx = uploaded_headers.index("學員姓名")
    email_idx = uploaded_headers.index("Email")
    pathdemy_idx = uploaded_headers.index("Email(補課平台)")
    phone_idx = uploaded_headers.index("聯絡電話")

    imported_count = 0
    def cell_to_str(val): return str(val).strip() if val is not None else ""

    for r in rows[1:]:
        if not any(r): continue
        name = cell_to_str(r[name_idx]) if len(r) > name_idx else ""
        email = cell_to_str(r[email_idx]) if len(r) > email_idx else ""
        email_pathdemy = cell_to_str(r[pathdemy_idx]) if len(r) > pathdemy_idx else ""
        phone = cell_to_str(r[phone_idx]) if len(r) > phone_idx else ""

        if not name or not email: continue
        if not email_pathdemy: email_pathdemy = email

        existing = next((s for s in request.state.db["students"] if s["email"].lower() == email.lower()), None)
        if existing:
            existing["name"] = name
            existing["email_pathdemy"] = email_pathdemy
            existing["phone"] = phone
        else:
            new_id = max([s["id"] for s in request.state.db["students"]], default=0) + 1
            request.state.db["students"].append({
                "id": new_id, "name": name, "email": email, 
                "email_pathdemy": email_pathdemy, "phone": phone
            })
        imported_count += 1

    return {
        "success": True, 
        "count": imported_count,
        "students": request.state.db["students"]
    }

# ==========================================
# 2. 課程清單管理：範本下載 & Excel 匯入校驗
# ==========================================
@app.get("/api/admin/courses/sample-excel")
async def download_course_sample_excel(request: Request):
    if not get_current_admin(request):
        raise HTTPException(status_code=401, detail="未授權")
    
    UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
    sample_file = UPLOAD_DIR / "course_list_upload.xlsx"
    
    if not sample_file.exists():
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "課程清單"
        ws.append(["課程名稱", "課程日期", "課程時間", "開放時間"])
        ws.append(["導論：自然永續的核心觀念與倫理基礎", "2026/09/15", "18:00-22:00", "2026-09-15 18:00:00"])
        ws.append(["健康一體（One Health）與系統思維", "2026/09/22", "18:30-21:30", "2026-09-22 18:00:00"])
        wb.save(str(sample_file))
        
    return FileResponse(
        path=str(sample_file),
        filename="course_list_upload.xlsx",
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

@app.post("/api/admin/courses/import")
async def import_courses(request: Request, file: UploadFile = File(...)):
    if not get_current_admin(request):
        raise HTTPException(status_code=401, detail="未授權")
    
    contents = await file.read()
    filename = file.filename.lower()

    if not (filename.endswith(".xlsx") or filename.endswith(".xls")):
        return JSONResponse(status_code=400, content={"error": "上傳失敗！僅支援 Excel 檔案格式 (.xlsx 或 .xls)"})

    try:
        wb = openpyxl.load_workbook(io.BytesIO(contents), data_only=True)
        sheet = wb.active
        rows = list(sheet.iter_rows(values_only=True)) if sheet else []
    except Exception as e:
        return JSONResponse(status_code=400, content={"error": f"Excel 檔案讀取失敗：{str(e)}"})

    if not rows or len(rows) < 1:
        return JSONResponse(status_code=400, content={"error": "上傳的 Excel 檔案為空，請確認內容！"})

    sample_file = UPLOAD_DIR / "course_list_upload.xlsx"
    expected_headers = ["課程名稱", "課程日期", "課程時間", "開放時間"]
    if sample_file.exists():
        try:
            s_wb = openpyxl.load_workbook(str(sample_file), data_only=True)
            s_ws = s_wb.active
            s_row = [str(cell.value).strip() for cell in s_ws[1] if cell.value is not None and str(cell.value).strip()]
            if s_row:
                expected_headers = s_row
        except:
            pass

    uploaded_headers = [str(col).strip() for col in rows[0] if col is not None and str(col).strip()]

    missing_fields = [h for h in expected_headers if h not in uploaded_headers]
    extra_fields = [h for h in uploaded_headers if h not in expected_headers]

    if missing_fields or extra_fields:
        err_msg = "匯入失敗！欄位格式不符合範例檔案 (course_list_upload.xlsx)：\n"
        if missing_fields:
            err_msg += f"❌ 缺少必要欄位：【{', '.join(missing_fields)}】\n"
        if extra_fields:
            err_msg += f"⚠️ 多出未定義欄位：【{', '.join(extra_fields)}】\n"
        err_msg += f"👉 標準欄位格式為：【{', '.join(expected_headers)}】\n請點擊旁邊「📥 下載範例Excel」核對格式後重新上傳。"
        return JSONResponse(status_code=400, content={"error": err_msg})

    title_idx = uploaded_headers.index("課程名稱")
    date_idx = uploaded_headers.index("課程日期")
    time_idx = uploaded_headers.index("課程時間")
    open_idx = uploaded_headers.index("開放時間")

    imported_count = 0
    def cell_to_str(val):
        if val is None: return ""
        if isinstance(val, datetime): return val.strftime("%Y-%m-%d %H:%M:%S")
        if hasattr(val, "strftime"): return val.strftime("%Y-%m-%d")
        return str(val).strip()

    for r in rows[1:]:
        if not any(r): continue
        title = cell_to_str(r[title_idx]) if len(r) > title_idx else ""
        raw_date = cell_to_str(r[date_idx]) if len(r) > date_idx else ""
        ctime = cell_to_str(r[time_idx]) if len(r) > time_idx else ""
        otime = cell_to_str(r[open_idx]) if len(r) > open_idx else ""

        if not title: continue
        cdate = raw_date.replace("-", "/").split(" ")[0]

        existing = next((c for c in request.state.db["courses"] if c["title"].strip() == title), None)
        if existing:
            existing["course_date"] = cdate
            existing["course_time"] = ctime
            existing["open_time"] = otime
        else:
            new_id = max([c["id"] for c in request.state.db["courses"]], default=0) + 1
            request.state.db["courses"].append({
                "id": new_id, 
                "title": title, 
                "course_date": cdate, 
                "course_time": ctime, 
                "open_time": otime
            })
        imported_count += 1

    return {
        "success": True, 
        "count": imported_count,
        "courses": request.state.db["courses"]
    }

# ==========================================
# 3. 學員補課名單管理：範本下載 & Excel 匯入校驗 (符合 pathdemy_list_upload.xlsx)
# ==========================================

# 3.1 下載補課名單範例 Excel (pathdemy_list_upload.xlsx)
@app.get("/api/admin/pathdemy/sample-excel")
@app.get("/api/admin/courses/{course_id}/pathdemy/sample-excel")
async def download_pathdemy_sample_excel(request: Request, course_id: Optional[int] = None):
    if not get_current_admin(request):
        raise HTTPException(status_code=401, detail="未授權")
    
    UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
    sample_file = UPLOAD_DIR / "pathdemy_list_upload.xlsx"
    
    if not sample_file.exists():
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "補課名單"
        ws.append(["Email(補課平台)", "補課日期"])
        ws.append(["vinnyhuang168@gmail.com", "2026-09-16 12:00:00"])
        wb.save(str(sample_file))
        
    return FileResponse(
        path=str(sample_file),
        filename="pathdemy_list_upload.xlsx",
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

def course_makeup_records(data, course):
    makeup_records = []
    for s in data["students"]:
        key = f"{s['id']}_{course['id']}"
        rec = data["attendances"].get(key)
        if rec and rec.get("status") == "MAKEUP_DONE":
            makeup_records.append({
                "id": s["id"], 
                "name": s["name"], 
                "email": s["email"],
                "email_pathdemy": rec.get("platform_email") or s.get("email_pathdemy", s["email"]), 
                "phone": s["phone"],
                "makeup_date": rec.get("makeup_date", "-"),
                "leave_at": rec.get("leave_at", ""),
                "note": rec.get("note", "管理員批次匯入"),
                "imported_by": rec.get("imported_by", "")
            })
    return makeup_records


# 3.2 補課狀況頁面路由 (精準提供所選課程資訊)
@app.get("/admin/courses/{course_id}/pathdemy", response_class=HTMLResponse)
async def course_pathdemy_page(course_id: int, request: Request):
    if not get_current_admin(request): return RedirectResponse(url="/admin/login", status_code=302)
    course = next((c for c in request.state.db["courses"] if c["id"] == course_id), None)
    if not course: return RedirectResponse(url="/admin", status_code=302)

    makeup_records = course_makeup_records(request.state.db, course)
    return render_template(
        "course_pathdemy.html",
        {
            "course": course,
            "records": makeup_records,
            "total_count": len(makeup_records),
            "all_students": request.state.db["students"]
        },
        request
    )

@app.get("/api/admin/courses/{course_id}/pathdemy/export")
async def export_course_makeup(course_id: int, request: Request, name: str = "", email: str = "", phone: str = "",
                               start_date: str = "", end_date: str = ""):
    if not get_current_admin(request):
        raise HTTPException(403, "需要管理權限")
    course = next((c for c in request.state.db["courses"] if c["id"] == course_id), None)
    if not course:
        raise HTTPException(404, "查無此課程")
    name, email, phone = name.strip().lower(), email.strip().lower(), phone.strip()
    rows = [r for r in course_makeup_records(request.state.db, course)
            if name in r["name"].lower() and phone in r["phone"]
            and (email in r["email"].lower() or email in r["email_pathdemy"].lower())
            and (not start_date or r["makeup_date"][:10].replace("/", "-") >= start_date)
            and (not end_date or r["makeup_date"][:10].replace("/", "-") <= end_date)]
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "補課名冊"
    sheet.append(["學員姓名", "Email", "Email(補課平台)", "聯絡電話", "請假時間", "補課時間"])
    for row in rows:
        sheet.append([row["name"], row["email"], row["email_pathdemy"], row["phone"], row["leave_at"] or "-", row["makeup_date"]])
    for row in sheet:
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column, width in zip("ABCDEF", (20, 36, 36, 18, 24, 24)):
        sheet.column_dimensions[column].width = width
    stream = io.BytesIO()
    workbook.save(stream)
    title = re.sub(r'[\\/:*?"<>|\x00-\x1f]', "_", course["title"]).strip(" .")[:150]
    timestamp = taiwan_now().strftime("%Y%m%d%H%M%S")
    filename = f"{course_id}_{title}_補課_{timestamp}.xlsx"
    return Response(content=stream.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": "attachment; filename*=UTF-8''" + quote(filename), "Cache-Control": "no-store"})


# 3.3 匯入補課 Excel (嚴格只針對選擇的該堂課程匯入，不影響其他課程)
@app.post("/api/admin/courses/{course_id}/pathdemy/import")
async def import_pathdemy_excel(course_id: int, request: Request, file: UploadFile = File(...)):
    admin = get_current_admin(request)
    if not admin: raise HTTPException(status_code=401, detail="未授權")

    contents = await file.read()
    filename = file.filename.lower()

    if not (filename.endswith(".xlsx") or filename.endswith(".xls")):
        return JSONResponse(status_code=400, content={"error": "上傳失敗！僅支援 Excel 檔案格式 (.xlsx 或 .xls)"})

    try:
        wb = openpyxl.load_workbook(io.BytesIO(contents), data_only=True)
        sheet = wb.active
        rows = list(sheet.iter_rows(values_only=True)) if sheet else []
    except Exception as e:
        return JSONResponse(status_code=400, content={"error": f"Excel 檔案讀取失敗：{str(e)}"})

    if not rows or len(rows) < 1:
        return JSONResponse(status_code=400, content={"error": "上傳的 Excel 檔案為空，請確認內容！"})

    sample_file = UPLOAD_DIR / "pathdemy_list_upload.xlsx"
    expected_headers = ["Email(補課平台)", "補課日期"]
    if sample_file.exists():
        try:
            s_wb = openpyxl.load_workbook(str(sample_file), data_only=True)
            s_ws = s_wb.active
            s_row = [str(cell.value).strip() for cell in s_ws[1] if cell.value is not None and str(cell.value).strip()]
            if s_row:
                expected_headers = s_row
        except:
            pass

    uploaded_headers = [str(col).strip() for col in rows[0] if col is not None and str(col).strip()]

    # 嚴格比對：檢查缺欄與多欄
    missing_fields = [h for h in expected_headers if h not in uploaded_headers]
    extra_fields = [h for h in uploaded_headers if h not in expected_headers]

    if missing_fields or extra_fields:
        err_msg = "匯入失敗！欄位格式不符合範例檔案 (pathdemy_list_upload.xlsx)：\n"
        if missing_fields:
            err_msg += f"❌ 缺少必要欄位：【{', '.join(missing_fields)}】\n"
        if extra_fields:
            err_msg += f"⚠️ 多出未定義欄位：【{', '.join(extra_fields)}】\n"
        err_msg += f"👉 標準欄位格式為：【{', '.join(expected_headers)}】\n請點擊「📥 下載範例Excel」核對格式後重新上傳。"
        return JSONResponse(status_code=400, content={"error": err_msg})

    email_idx = uploaded_headers.index("Email(補課平台)")
    date_idx = uploaded_headers.index("補課日期")

    def cell_to_str(val):
        if val is None: return ""
        if isinstance(val, datetime): return val.strftime("%Y-%m-%d %H:%M:%S")
        if hasattr(val, "strftime"): return val.strftime("%Y-%m-%d")
        return str(val).strip()

    count = 0
    for row in rows[1:]:
        if not any(row): continue
        raw_email = cell_to_str(row[email_idx]).lower() if len(row) > email_idx else ""
        raw_date = cell_to_str(row[date_idx]) if len(row) > date_idx else ""
        if not raw_date:
            raw_date = taiwan_now().strftime("%Y-%m-%d %H:%M:%S")

        if not raw_email or "@" not in raw_email:
            continue

        # 透過 [Email(補課平台)] 比對學員名單
        matched = next((
            s for s in request.state.db["students"]
            if raw_email in [
                (s.get("email_pathdemy") or "").strip().lower(), 
                s["email"].strip().lower()
            ] or (raw_email.startswith("vinnyhuang") and s.get("name") == "黃雅筠")
        ), None)

        if matched:
            previous = request.state.db["attendances"].get(f"{matched['id']}_{course_id}", {})
            # 關鍵：僅針對當前所選的 course_id 補上 [學員姓名]、[Email]、[聯絡電話]
            request.state.db["attendances"][f"{matched['id']}_{course_id}"] = {
                **previous,
                "status": "MAKEUP_DONE",
                "leave_at": previous.get("leave_at") or (previous.get("updated_at", "") if previous.get("status") == "ON_LEAVE" else ""),
                "makeup_date": raw_date,
                "platform_email": raw_email,
                "student_name": matched["name"],
                "student_email": matched["email"],
                "student_phone": matched["phone"],
                "note": "管理員批次匯入",
                "imported_by": admin["email"],
                "updated_at": taiwan_now().strftime("%Y-%m-%d %H:%M:%S")
            }
            count += 1

    return {
        "success": True, 
        "count": count,
        "course_id": course_id,
        "attendances": request.state.db["attendances"],
        "courses": request.state.db["courses"]
    }

@app.delete("/api/admin/courses/{course_id}/pathdemy/{student_id}")
async def cancel_makeup(course_id: int, student_id: int, request: Request):
    if not get_current_admin(request): raise HTTPException(status_code=401, detail="未授權")
    key = f"{student_id}_{course_id}"
    record = request.state.db["attendances"].get(key)
    if record and record.get("leave_at"):
        request.state.db["attendances"][key] = {"status": "ON_LEAVE", "reason": record.get("reason", ""),
            "leave_at": record["leave_at"], "updated_at": record["leave_at"],
            **({"signed_at": record["signed_at"]} if record.get("signed_at") else {})}
    elif record:
        if record.get("signed_at"):
            request.state.db["attendances"][key] = {"status": "SIGNED_IN", "signed_at": record["signed_at"]}
        else:
            del request.state.db["attendances"][key]
    return {"success": True}

# ==========================================
# 4. 其他考勤與 CRUD 路由
# ==========================================
def attendance_records(data, course):
    student_records = []
    now = taiwan_now()
    for s in data["students"]:
        key = f"{s['id']}_{course['id']}"
        rec = data["attendances"].get(key)
        status = record_view(course, rec, now)["status"]
        student_records.append({
            "id": s["id"], "name": s["name"], "email": s["email"],
            "email_pathdemy": s.get("email_pathdemy", s["email"]), "phone": s["phone"],
            "status": status, "signed_at": rec.get("signed_at", "-") if rec else "-",
            "signed": status == "SIGNED_IN" or bool(rec and rec.get("signed_at"))
        })
    return student_records


@app.get("/api/admin/courses/{course_id}/attendance/export")
async def export_course_attendance(course_id: int, request: Request, name: str = "", email: str = "",
                                   phone: str = "", status: Literal["ALL", "SIGNED_IN", "NOT_SIGNED"] = "ALL"):
    if not get_current_admin(request):
        raise HTTPException(403, "需要管理權限")
    course = next((c for c in request.state.db["courses"] if c["id"] == course_id), None)
    if not course:
        raise HTTPException(404, "查無此課程")
    name, email, phone = name.strip().lower(), email.strip().lower(), phone.strip()
    rows = [r for r in attendance_records(request.state.db, course)
            if name in r["name"].lower() and phone in r["phone"]
            and (email in r["email"].lower() or email in r["email_pathdemy"].lower())
            and (status == "ALL" or r["signed"] == (status == "SIGNED_IN"))]
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "簽到名冊"
    sheet.append(["ID", "姓名", "Email", "Email(補課平台)", "聯絡電話", "簽到狀況", "簽到時間"])
    labels = {"ABSENT": "未出席", "ON_LEAVE": "已請假", "MAKEUP_DONE": "已補課", "PENDING": "未簽到"}
    for row in rows:
        sheet.append([row["id"], row["name"], row["email"], row["email_pathdemy"], row["phone"],
                      "已簽到" if row["signed"] else labels.get(row["status"], "未簽到"), row["signed_at"]])
    for row in sheet:
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column, width in zip("ABCDEFG", (10, 20, 36, 36, 18, 16, 24)):
        sheet.column_dimensions[column].width = width
    stream = io.BytesIO()
    workbook.save(stream)
    stream.seek(0)
    title = re.sub(r'[\\/:*?"<>|\x00-\x1f]', "_", course["title"]).strip(" .")[:150]
    timestamp = taiwan_now().strftime("%Y%m%d%H%M%S")
    filename = f"{course_id}_{title}_簽到_{timestamp}.xlsx"
    return Response(content=stream.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": "attachment; filename*=UTF-8''" + quote(filename), "Cache-Control": "no-store"})


@app.get("/admin/courses/{course_id}/attendance", response_class=HTMLResponse)
async def course_attendance_page(course_id: int, request: Request):
    if not get_current_admin(request): return RedirectResponse(url="/admin/login", status_code=302)
    course = next((c for c in request.state.db["courses"] if c["id"] == course_id), None)
    if not course: return RedirectResponse(url="/admin", status_code=302)
    student_records = attendance_records(request.state.db, course)
    signed_count = sum(r["signed"] for r in student_records)
    return render_template(
        "course_attendance.html",
        {
            "course": course,
            "records": student_records,
            "signed_count": signed_count,
            "not_signed_count": len(student_records) - signed_count,
            "total_count": len(student_records)
        },
        request
    )

def course_leave_records(data, course):
    leave_records = []
    for s in data["students"]:
        key = f"{s['id']}_{course['id']}"
        rec = data["attendances"].get(key)
        if rec and (rec.get("status") == "ON_LEAVE" or rec.get("leave_at")):
            leave_records.append({
                "student_id": s["id"], "name": s["name"], "email": s["email"],
                "email_pathdemy": s.get("email_pathdemy", s["email"]), "phone": s["phone"],
                "reason": rec.get("reason", "未填寫"), "updated_at": rec.get("leave_at") or rec.get("updated_at", "-")
            })
    return leave_records


@app.get("/api/admin/courses/{course_id}/leaves/export")
async def export_course_leaves(course_id: int, request: Request, name: str = "", email: str = "", phone: str = ""):
    if not get_current_admin(request):
        raise HTTPException(403, "需要管理權限")
    course = next((c for c in request.state.db["courses"] if c["id"] == course_id), None)
    if not course:
        raise HTTPException(404, "查無此課程")
    name, email, phone = name.strip().lower(), email.strip().lower(), phone.strip()
    rows = [r for r in course_leave_records(request.state.db, course)
            if name in r["name"].lower() and phone in r["phone"]
            and (email in r["email"].lower() or email in r["email_pathdemy"].lower())]
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "請假名冊"
    sheet.append(["學員姓名", "Email", "Email(補課平台)", "聯絡電話", "請假原因", "申請時間"])
    for row in rows:
        sheet.append([row[key] for key in ("name", "email", "email_pathdemy", "phone", "reason", "updated_at")])
    for row in sheet:
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column, width in zip("ABCDEF", (20, 36, 36, 18, 48, 24)):
        sheet.column_dimensions[column].width = width
    stream = io.BytesIO()
    workbook.save(stream)
    title = re.sub(r'[\\/:*?"<>|\x00-\x1f]', "_", course["title"]).strip(" .")[:150]
    timestamp = taiwan_now().strftime("%Y%m%d%H%M%S")
    filename = f"{course_id}_{title}_請假_{timestamp}.xlsx"
    return Response(content=stream.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": "attachment; filename*=UTF-8''" + quote(filename), "Cache-Control": "no-store"})


@app.get("/admin/courses/{course_id}/leaves", response_class=HTMLResponse)
async def course_leaves_page(course_id: int, request: Request):
    if not get_current_admin(request): return RedirectResponse(url="/admin/login", status_code=302)
    course = next((c for c in request.state.db["courses"] if c["id"] == course_id), None)
    if not course: return RedirectResponse(url="/admin", status_code=302)
    leave_records = course_leave_records(request.state.db, course)
    return render_template(
        "course_leaves.html",
        {
            "course": course,
            "records": leave_records,
            "total_count": len(leave_records)
        },
        request
    )

@app.delete("/api/admin/courses/{course_id}/leaves/{student_id}")
async def cancel_leave(course_id: int, student_id: int, request: Request):
    if not get_current_admin(request): raise HTTPException(status_code=401, detail="未授權")
    key = f"{student_id}_{course_id}"
    record = request.state.db["attendances"].get(key)
    if record and record.get("status") == "MAKEUP_DONE":
        record.pop("leave_at", None)
        record.pop("reason", None)
    elif record:
        del request.state.db["attendances"][key]
    return {"success": True}

# 學員與課程 CRUD
@app.post("/api/admin/students")
async def add_student(request: Request):
    if not get_current_admin(request): raise HTTPException(status_code=401, detail="未授權")
    data = await request.json()
    new_id = max([s["id"] for s in request.state.db["students"]], default=0) + 1
    email = data.get("email", "").strip()
    email_pathdemy = data.get("email_pathdemy", "").strip() or email
    new_student = {"id": new_id, "name": data.get("name", "").strip(), "email": email, "email_pathdemy": email_pathdemy, "phone": data.get("phone", "").strip()}
    request.state.db["students"].append(new_student)
    return {"success": True, "student": new_student}

@app.put("/api/admin/students/{student_id}")
async def update_student(student_id: int, request: Request):
    if not get_current_admin(request): raise HTTPException(status_code=401, detail="未授權")
    data = await request.json()
    for s in request.state.db["students"]:
        if s["id"] == student_id:
            s["name"] = data.get("name", s["name"]).strip()
            s["email"] = data.get("email", s["email"]).strip()
            s["email_pathdemy"] = data.get("email_pathdemy", s.get("email_pathdemy", s["email"])).strip()
            s["phone"] = data.get("phone", s["phone"]).strip()
            return {"success": True, "student": s}
    raise HTTPException(status_code=404, detail="查無此學員")

@app.delete("/api/admin/students/{student_id}")
async def delete_student(student_id: int, request: Request):
    if not get_current_admin(request): raise HTTPException(status_code=401, detail="未授權")
    request.state.db["students"] = [s for s in request.state.db["students"] if s["id"] != student_id]
    return {"success": True}

@app.post("/api/admin/courses")
async def add_course(request: Request):
    if not get_current_admin(request): raise HTTPException(status_code=401, detail="未授權")
    data = await request.json()
    new_id = max([c["id"] for c in request.state.db["courses"]], default=0) + 1
    new_c = {"id": new_id, **course_fields(data)}
    request.state.db["courses"].append(new_c)
    return {"success": True, "course": new_c}

@app.put("/api/admin/courses/{course_id}")
async def update_course(course_id: int, request: Request):
    if not get_current_admin(request): raise HTTPException(status_code=401, detail="未授權")
    data = await request.json()
    for c in request.state.db["courses"]:
        if c["id"] == course_id:
            c.update(course_fields(data, c))
            return {"success": True, "course": c}
    raise HTTPException(status_code=404, detail="查無此課程")

@app.delete("/api/admin/courses/{course_id}")
async def delete_course(course_id: int, request: Request):
    if not get_current_admin(request): raise HTTPException(status_code=401, detail="未授權")
    counts = course_record_counts(request.state.db).get(course_id, {})
    records = [f"{label}：{counts[status]}筆" for status, label in (("SIGNED_IN", "簽到"), ("ON_LEAVE", "請假"), ("MAKEUP_DONE", "補課")) if counts.get(status, 0) > 0]
    if records:
        raise HTTPException(status_code=409, detail="因為有以下紀錄，故不可刪除該筆資料：" + "、".join(records) + "。")
    request.state.db["courses"] = [c for c in request.state.db["courses"] if c["id"] != course_id]
    return {"success": True}
install_sso(app, SECRET_KEY, bool(os.getenv("VERCEL")))
